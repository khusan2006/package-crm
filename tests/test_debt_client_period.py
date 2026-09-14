"""Qarz — bitta mijoz sahifasidagi oraliq sverka.

The list page answers "which of my clients took what this month"; this page answers it
for one client and shows the movements behind the answer, so the figures can be checked
line by line with the client sitting opposite you.

Pinned down here: the window is opt-in, the figures match the list page's row for the
same client, the movements are cut to the window while the running balance still counts
from the beginning, and the Excel gains the sverka sheet only when a window is picked.
"""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from crm.models import Client, Payment, Product, Sale, SaleItem

pytestmark = pytest.mark.django_db

TODAY = timezone.localdate()


@pytest.fixture
def product(db):
    return Product.objects.create(
        name="Sverka paket", sku="SVR-002",
        cost_price=Decimal("10000"), price=Decimal("15000"),
    )


def _sale(client, seller, *, day, kg, price, product, paid=None):
    sale = Sale.objects.create(
        client=client, sales_rep=seller, date=day,
        debt_deadline=day + timedelta(days=7), debt_term_days=7,
    )
    SaleItem.objects.create(
        sale=sale, product=product, dimension=Sale.Dimension.KG,
        weight=Decimal(kg), price=Decimal(price), cost_price=Decimal("10000"),
    )
    if paid:
        Payment.objects.create(
            sale=sale, amount=Decimal(paid), amount_original=Decimal(paid),
            method=Payment.Method.CASH, kind=Payment.Kind.SALE,
            date=day, created_by=seller,
        )
    return sale


@pytest.fixture
def traded(db, seller_user, product):
    """One client, three receipts: one before the window and two inside it."""
    mijoz = Client.objects.create(name="Sverka mijoz", owner=seller_user)
    dan, gacha = TODAY - timedelta(days=14), TODAY
    _sale(mijoz, seller_user, day=dan - timedelta(days=20), kg="5", price="15000",
          product=product)
    _sale(mijoz, seller_user, day=dan + timedelta(days=1), kg="10", price="15000",
          product=product, paid="60000")
    _sale(mijoz, seller_user, day=gacha, kg="4", price="20000", product=product)
    return {"client": mijoz, "dan": dan, "gacha": gacha}


def _window(traded):
    return {"dan": traded["dan"].isoformat(), "gacha": traded["gacha"].isoformat()}


def test_no_window_keeps_the_page_as_it_was(client, admin_user, traded):
    client.force_login(admin_user)
    resp = client.get(f"/debts/{traded['client'].pk}/")

    assert resp.status_code == 200
    assert resp.context["period_window"] is False
    assert resp.context["period"] is None
    assert resp.context["movements"] == []
    # The picker is offered, so the window can be picked from here.
    body = resp.content.decode()
    assert "Oraliq sverka" in body and "daterange" in body


def test_window_figures_match_the_list_page_row(client, admin_user, traded):
    """One client's own page and their row in the Qarzlar list are the same numbers —
    if these two ever drift, one of the screens is lying."""
    client.force_login(admin_user)
    mijoz = traded["client"]
    detail = client.get(f"/debts/{mijoz.pk}/", _window(traded)).context["period"]
    listed = next(
        r for r in client.get("/debts/", _window(traded)).context["debtors"]
        if r["client"] == mijoz
    )["period"]

    assert detail["net_taken"] == listed["net_taken"] == Decimal("230000")
    assert detail["net_taken_kg"] == listed["net_taken_kg"] == Decimal("14")
    assert detail["paid"] == listed["paid"] == Decimal("60000")
    # Took 230 000, paid 60 000 — the debt grew by the difference over the window.
    assert detail["net_change"] == Decimal("170000")


def test_movements_are_cut_to_the_window_but_the_balance_is_not(
    client, admin_user, traded
):
    """Only what happened inside the window is listed; the running balance still counts
    from the client's first ever receipt, because that is what they actually owe."""
    client.force_login(admin_user)
    resp = client.get(f"/debts/{traded['client'].pk}/", _window(traded))
    movements = resp.context["movements"]

    assert all(traded["dan"] <= m["date"] <= traded["gacha"] for m in movements)
    # 2 sales + 1 payment inside; the pre-window sale is not listed…
    assert len(movements) == 3
    # …but its 75 000 is still carried in the balance the rows show.
    assert movements[-1]["balance"] == Decimal("245000")
    assert resp.context["total"] == Decimal("245000")


def test_excel_gains_the_sverka_sheet_only_with_a_window(client, admin_user, traded):
    client.force_login(admin_user)
    mijoz = traded["client"]

    plain = load_workbook(BytesIO(client.get(f"/debts/{mijoz.pk}/export/").content))
    assert plain.sheetnames == ["Ochiq cheklar"]

    book = load_workbook(
        BytesIO(client.get(f"/debts/{mijoz.pk}/export/", _window(traded)).content)
    )
    span = f"{traded['dan'].strftime('%d.%m.%Y')}–{traded['gacha'].strftime('%d.%m.%Y')}"
    assert book.sheetnames == ["Ochiq cheklar", f"Oraliq {span}"]

    sheet = book[f"Oraliq {span}"]
    figures = {r[1].value: r[3].value for r in sheet.iter_rows(min_row=2, max_row=5)}
    assert figures["Olgan yuk"] == 230000
    assert figures["To'lagan pul"] == 60000
    assert figures["Qarz o'zgarishi"] == 170000
    # The movements follow the summary, under their own heading.
    labels = [r[1].value for r in sheet.iter_rows(min_row=6)]
    assert labels.count("Sotuv") == 2


def test_seller_sees_only_their_own_movements(
    django_user_model, client, seller_user, traded, product
):
    other = django_user_model.objects.create_user(
        username="boshqa_sotuvchi2", password="x", role=django_user_model.Role.SALES,
    )
    _sale(traded["client"], other, day=traded["gacha"], kg="9", price="15000",
          product=product, paid="100000")

    client.force_login(seller_user)
    resp = client.get(f"/debts/{traded['client'].pk}/", _window(traded))

    assert resp.context["period"]["net_taken"] == Decimal("230000")
    assert resp.context["period"]["paid"] == Decimal("60000")
    assert all(m["user"] == seller_user for m in resp.context["movements"])
