"""Qarz muddati — a client who pays is not late since the day the debt was written.

The case the business asked for, word for word: a client has owed 20 500 000 for 80
days, pays 20 000 000 on some day, and the 500 000 left is counted as due from that
day — on the Qarzlar page and in every Excel that is downloaded from it.

It is the CLIENT's payment that restarts the counter, not the receipt's. Money goes
to the oldest receipt first, so the part left over usually sits on a receipt the
payment never touched; that one restarts all the same.
"""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from crm.models import (
    Client,
    Payment,
    Product,
    Sale,
    SaleItem,
    recompute_client_debt_deadlines,
)

pytestmark = pytest.mark.django_db

TODAY = timezone.localdate()
PAID_ON = TODAY - timedelta(days=10)


def _receipt(client, seller, product, *, days_ago, amount, term=7):
    day = TODAY - timedelta(days=days_ago)
    sale = Sale.objects.create(
        client=client, sales_rep=seller, date=day,
        debt_deadline=day + timedelta(days=term), debt_term_days=term,
    )
    SaleItem.objects.create(
        sale=sale, product=product, dimension=Sale.Dimension.KG,
        weight=Decimal("1000"), price=Decimal(amount) / 1000, cost_price=Decimal("10"),
    )
    return sale


@pytest.fixture
def mijoz(seller_user):
    return Client.objects.create(name="Qarzdor mijoz", owner=seller_user)


@pytest.fixture
def product(db):
    return Product.objects.create(
        name="Muddat paket", sku="MDT-001",
        cost_price=Decimal("10"), price=Decimal("15"),
    )


def _pay(client, seller, mijoz, amount, day=PAID_ON):
    client.force_login(seller)
    client.post(
        f"/debts/{mijoz.pk}/pay/",
        {"amount": amount, "method": "cash", "date": day.isoformat()},
    )


def _column(rows, name):
    return [row[rows[0].index(name)] for row in rows[1:]]


def _xlsx(client, url):
    book = load_workbook(BytesIO(client.get(url).content))
    return [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]


def test_80_days_late_pays_20_million_and_the_rest_counts_from_that_day(
    client, admin_user, seller_user, mijoz, product
):
    sale = _receipt(mijoz, seller_user, product, days_ago=80, amount="20500000")
    assert sale.is_overdue

    _pay(client, seller_user, mijoz, "20000000")

    sale.refresh_from_db()
    assert sale.debt_remaining == Decimal("500000")
    assert sale.debt_deadline == PAID_ON

    client.force_login(admin_user)
    row = next(
        r for r in client.get("/debts/").context["debtors"] if r["client"] == mijoz
    )
    assert row["earliest"] == PAID_ON

    debtors = _xlsx(client, "/debts/export/")
    assert _column(debtors, "Kechikkan kun") == [10]  # not 73
    assert _column(debtors, "Eng yaqin muddat") == [PAID_ON.strftime("%d.%m.%Y")]

    receipts = _xlsx(client, f"/debts/{mijoz.pk}/export/")
    assert _column(receipts, "Kechikkan kun") == [10]


def test_the_rest_left_on_a_receipt_the_money_never_reached_restarts_too(
    client, admin_user, seller_user, mijoz, product
):
    # 20 000 000 clears the older receipt exactly; the 500 000 is the newer one,
    # which no part of the payment was booked to.
    older = _receipt(mijoz, seller_user, product, days_ago=80, amount="20000000")
    newer = _receipt(mijoz, seller_user, product, days_ago=60, amount="500000")

    _pay(client, seller_user, mijoz, "20000000")

    older.refresh_from_db()
    newer.refresh_from_db()
    assert older.is_paid
    assert not newer.payments.exists()
    assert newer.debt_deadline == PAID_ON

    client.force_login(admin_user)
    assert _column(_xlsx(client, "/debts/export/"), "Kechikkan kun") == [10]  # not 53
    assert _column(_xlsx(client, f"/debts/{mijoz.pk}/export/"), "Kechikkan kun") == [10]


def test_goods_taken_after_the_payment_keep_their_own_term(
    client, seller_user, mijoz, product
):
    _receipt(mijoz, seller_user, product, days_ago=80, amount="20500000")
    fresh = _receipt(mijoz, seller_user, product, days_ago=5, amount="3000000", term=14)

    _pay(client, seller_user, mijoz, "20000000")

    fresh.refresh_from_db()
    assert fresh.debt_deadline == fresh.date + timedelta(days=14)
    assert not fresh.is_overdue


def test_paying_part_before_the_deadline_does_not_make_the_client_late_sooner(
    client, seller_user, mijoz, product
):
    sale = _receipt(mijoz, seller_user, product, days_ago=3, amount="2000000", term=14)

    _pay(client, seller_user, mijoz, "500000", day=TODAY)

    sale.refresh_from_db()
    assert sale.debt_deadline == sale.date + timedelta(days=14)


def test_voiding_the_payment_puts_every_receipt_back(
    client, admin_user, seller_user, mijoz, product
):
    older = _receipt(mijoz, seller_user, product, days_ago=80, amount="20000000")
    newer = _receipt(mijoz, seller_user, product, days_ago=60, amount="500000")
    _pay(client, seller_user, mijoz, "20000000")

    payment = Payment.objects.get(sale=older)
    client.force_login(admin_user)
    client.post(f"/payments/{payment.pk}/delete/")

    assert not Payment.objects.filter(pk=payment.pk).exists()
    newer.refresh_from_db()
    assert newer.debt_deadline == newer.date + timedelta(days=7)


def test_money_taken_at_the_counter_restarts_nothing(
    client, seller_user, mijoz, product
):
    # A new purchase part-paid on the day is that sale's own money, not a repayment
    # of the old debt — the old receipt keeps counting from its own deadline.
    old = _receipt(mijoz, seller_user, product, days_ago=40, amount="1000000")
    new = _receipt(mijoz, seller_user, product, days_ago=2, amount="800000")
    Payment.objects.create(
        sale=new, amount=Decimal("300000"), amount_original=Decimal("300000"),
        method=Payment.Method.CASH, kind=Payment.Kind.SALE,
        date=new.date, created_by=seller_user,
    )
    recompute_client_debt_deadlines(mijoz.pk)

    old.refresh_from_db()
    assert old.debt_deadline == old.date + timedelta(days=7)
