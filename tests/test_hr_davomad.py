"""HR — davomad, dam va bayram, va ulardan chiqadigan oylik.

The wage a month adds to a worker's balance is no longer the salary: it is the days
on the davomad sheet times the day's pay. These tests pin that arithmetic down, the
way it feeds the month-to-month balance that was already there, and the rule that
marking a day is the supervisor's job — a seller reads the calendar but cannot move
a wage.

July 2026 is the month most of them work in: 31 days, four Sundays, 27 working days.
A month in the PAST is fixed on purpose — the current month stops counting at today,
which is right for the app and useless for arithmetic you want to check by hand.
"""

from datetime import date, time
from decimal import Decimal
from io import BytesIO

import pytest
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from crm.forms import EmployeeForm
from crm.models import (
    Attendance,
    DayHours,
    Employee,
    Expense,
    Holiday,
    Leave,
    SalaryRate,
)

pytestmark = pytest.mark.django_db

YEAR, MONTH = 2026, 7
WAGE = Decimal("2700000")          # 27 working days in July -> 100 000 a day
DAY = Decimal("100000")
MONDAY = date(2026, 7, 6)
SUNDAY = date(2026, 7, 5)


def make_worker(name="Ишчи Аваз", salary=WAGE, start=date(2026, 7, 1)):
    return Employee.objects.create(name=name, salary=salary, start_month=start)


def mark(worker, day, present=True):
    return Attendance.objects.create(employee=worker, date=day, present=present)


@pytest.fixture(autouse=True)
def sheet_on_from_july(settings):
    """The davomad sheet is switched on from July 2026 here — the month these tests do
    their arithmetic in — whichever month the deployment itself went live in."""
    settings.DAVOMAD_START_DATE = date(2026, 7, 1)


@pytest.fixture
def on_august_12(monkeypatch):
    """Wednesday 12 August 2026 — a current month that is partly over, with a day
    still open on the sheet (the 11th) and one already closed (the 10th)."""
    today = date(2026, 8, 12)
    monkeypatch.setattr(timezone, "localdate", lambda *args, **kwargs: today)
    return today


# --- what a month is worth ---------------------------------------------------------


def test_a_full_month_is_exactly_the_salary():
    worker = make_worker()
    assert worker.working_days_in(YEAR, MONTH) == 27
    assert worker.daily_rate_in(YEAR, MONTH) == DAY
    assert worker.earned_in(YEAR, MONTH) == WAGE


def test_a_salary_that_does_not_divide_still_comes_out_exact():
    """3 000 000 over 27 days is 111 111.11… a day; rounding the day first would put
    a full month a few so'm off the salary it is meant to reproduce."""
    worker = make_worker(salary=Decimal("3000000"))
    assert worker.earned_in(YEAR, MONTH) == Decimal("3000000")


def test_a_missed_day_comes_off():
    worker = make_worker()
    mark(worker, MONDAY, present=False)
    assert worker.earned_in(YEAR, MONTH) == WAGE - DAY
    assert worker.days_in(YEAR, MONTH) == (26, 0)


def test_a_sunday_worked_goes_on_top():
    worker = make_worker()
    mark(worker, SUNDAY)
    assert worker.earned_in(YEAR, MONTH) == WAGE + DAY


def test_a_day_is_worked_or_it_is_not(client, admin_user, on_august_12):
    """There is no half: the cell and the day card offer came / did not, and anything
    else a form sends is left alone rather than guessed at."""
    worker = make_worker()
    client.force_login(admin_user)
    for junk in ("plus", "yarim"):
        post_sheet(client, worker, {("d", "2026-08-11"): junk})
        client.post(
            reverse("employee_day", args=[worker.pk, "2026-08-11"]),
            {"action": junk, "note": ""},
        )
    assert not Attendance.objects.exists()
    sheet = client.get(reverse("attendance_grid"), {"oy": "2026-08"})
    cycles = {c["cycle"] for r in sheet.context["rows"] for c in r["cells"]}
    assert cycles == {"keldi,kelmadi", "dam,keldi"}


def test_a_closed_holiday_costs_nobody_anything():
    """The day leaves the divisor, so the month is worth the same salary over fewer
    days — and coming in on it is extra, like a Sunday."""
    worker = make_worker()
    Holiday.objects.create(date=date(2026, 7, 15), name="Firma kuni", rest=True)
    assert worker.working_days_in(YEAR, MONTH) == 26
    assert worker.earned_in(YEAR, MONTH) == WAGE
    mark(worker, date(2026, 7, 15))
    assert worker.earned_in(YEAR, MONTH) == (WAGE * 27 / 26).quantize(Decimal("0.01"))


def test_an_open_holiday_is_only_a_label():
    worker = make_worker()
    Holiday.objects.create(date=date(2026, 7, 15), name="Firma kuni", rest=False)
    assert worker.working_days_in(YEAR, MONTH) == 27
    mark(worker, date(2026, 7, 15), present=False)
    assert worker.earned_in(YEAR, MONTH) == WAGE - DAY


def test_paid_leave_keeps_the_wage_and_unpaid_leave_does_not():
    paid = make_worker(name="Haqli")
    Leave.objects.create(employee=paid, date=MONDAY, paid=True)
    assert paid.earned_in(YEAR, MONTH) == WAGE

    unpaid = make_worker(name="O'z hisobidan")
    Leave.objects.create(employee=unpaid, date=MONDAY, paid=False)
    assert unpaid.earned_in(YEAR, MONTH) == WAGE - DAY
    # Leave is a plan; turning up anyway is what happened, and it wins.
    mark(unpaid, MONDAY)
    assert unpaid.earned_in(YEAR, MONTH) == WAGE


def test_the_current_month_counts_only_the_days_that_have_come(on_august_12):
    """August 2026 has 26 working days; by the 12th ten of them are behind us."""
    worker = make_worker(salary=Decimal("2600000"))
    assert worker.earned_in(2026, 8) == Decimal("1000000")
    assert worker.earned_in(2026, 9) == Decimal("0")       # nobody has turned up yet


def test_a_raise_does_not_reprice_the_days_of_an_earlier_month(admin_user):
    worker = make_worker()
    SalaryRate.objects.create(employee=worker, effective_from=date(2026, 7, 1), amount=WAGE)
    SalaryRate.objects.create(
        employee=worker, effective_from=date(2026, 8, 1), amount=Decimal("5200000")
    )
    worker.salary = Decimal("5200000")
    worker.save()
    mark(worker, MONDAY, present=False)
    # July's missed day is priced at July's wage, not at the raise agreed since.
    assert worker.earned_in(YEAR, MONTH) == WAGE - DAY


# --- the balance that carries on -----------------------------------------------------


def test_the_balance_carries_what_the_sheet_says_not_the_salary(on_august_12, seller_user):
    worker = make_worker(salary=Decimal("2700000"))
    mark(worker, MONDAY, present=False)
    Expense.objects.create(
        date=date(2026, 7, 20), amount=Decimal("1000000"), category="Oylik / xodim",
        method="cash", employee=worker, created_by=seller_user,
    )
    # July: 26 days worked, 1 000 000 drawn.
    assert worker.balance_through(YEAR, MONTH) == Decimal("1600000")


def test_the_payroll_page_agrees_with_the_model(client, on_august_12, payroll_seller):
    worker = make_worker()
    mark(worker, MONDAY, present=False)
    client.force_login(payroll_seller)

    july = client.get(reverse("employee_list"), {"oy": "2026-07"})
    row = {r["employee"].pk: r for r in july.context["rows"]}[worker.pk]
    assert row["salary"] == WAGE
    assert row["earned"] == WAGE - DAY
    assert (row["days"], row["missed_days"], row["working_days"]) == (26, 1, 27)
    assert row["remaining"] == worker.balance_through(YEAR, MONTH)

    august = client.get(reverse("employee_list"), {"oy": "2026-08"})
    row = {r["employee"].pk: r for r in august.context["rows"]}[worker.pk]
    assert row["carried"] == WAGE - DAY                     # July rides in as worked
    assert row["remaining"] == worker.balance_through(2026, 8)


def test_a_finished_month_nobody_marked_carries_as_the_salary(client, on_august_12, payroll_seller):
    worker = make_worker()
    client.force_login(payroll_seller)
    august = client.get(reverse("employee_list"), {"oy": "2026-08"})
    row = {r["employee"].pk: r for r in august.context["rows"]}[worker.pk]
    assert row["carried"] == WAGE


# --- the davomad sheet -----------------------------------------------------------------


def post_sheet(client, worker, cells, month="2026-08"):
    data = {f"{kind}-{worker.pk}-{day}": value for (kind, day), value in cells.items()}
    return client.post(f"{reverse('attendance_grid')}?oy={month}", data)


# --- before the sheet was switched on ----------------------------------------------------
#
# June 2026 is the month before the sheet here. It was paid as the flat salary, and
# nothing written since may re-price it.

JUNE_MONDAY = date(2026, 6, 8)


def test_a_month_before_the_sheet_is_the_flat_salary_whatever_is_written():
    worker = make_worker(start=date(2026, 6, 1))
    mark(worker, JUNE_MONDAY, present=False)
    Leave.objects.create(employee=worker, date=date(2026, 6, 9), paid=False)
    Holiday.objects.create(date=date(2026, 6, 10), name="Firma kuni", rest=True)
    mark(worker, date(2026, 6, 10))
    assert worker.earned_in(2026, 6) == WAGE
    assert worker.balance_through(2026, 6) == WAGE


def test_the_payroll_shows_no_days_for_a_month_before_the_sheet(
    client, on_august_12, payroll_seller
):
    worker = make_worker(start=date(2026, 6, 1))
    mark(worker, JUNE_MONDAY, present=False)
    client.force_login(payroll_seller)

    june = client.get(reverse("employee_list"), {"oy": "2026-06"})
    row = {r["employee"].pk: r for r in june.context["rows"]}[worker.pk]
    assert (row["earned"], row["by_day"]) == (WAGE, False)
    assert "days" not in row
    assert june.context["by_day"] is False

    # June rides into July whole, despite the mark sitting against it.
    july = client.get(reverse("employee_list"), {"oy": "2026-07"})
    row = {r["employee"].pk: r for r in july.context["rows"]}[worker.pk]
    assert (row["carried"], row["by_day"]) == (WAGE, True)

    # The worker's own page says so rather than drawing a calendar nobody filled in.
    detail = client.get(reverse("employee_detail", args=[worker.pk]), {"oy": "2026-06"})
    assert detail.context["by_day"] is False
    assert "sched-month" not in detail.content.decode()


def test_the_sheet_and_the_day_card_do_not_open_before_the_first_month(
    client, admin_user, on_august_12
):
    worker = make_worker(start=date(2026, 6, 1))
    client.force_login(admin_user)

    for name in ("attendance_grid", "attendance_excel"):
        response = client.get(reverse(name), {"oy": "2026-06"})
        assert response.status_code == 302
        assert response.url == f"{reverse(name)}?oy=2026-07"

    # Not even a cell opened on purpose — there is no sheet for that month to open.
    posted = post_sheet(
        client, worker,
        {("d", "2026-06-08"): "kelmadi", ("o", "2026-06-08"): "1"}, month="2026-06",
    )
    assert posted.status_code == 302
    card = reverse("employee_day", args=[worker.pk, "2026-06-08"])
    assert client.get(card).status_code == 404
    assert client.post(card, {"action": "kelmadi", "note": ""}).status_code == 404
    assert client.post(card, {"action": "dam_haqli", "note": ""}).status_code == 404
    assert not Attendance.objects.exists()
    assert not Leave.objects.exists()

    # The first month offers no way back, and the payroll's link lands on it.
    first = client.get(reverse("attendance_grid"), {"oy": "2026-07"})
    assert first.context["prev_month"] == ""
    assert [o["value"] for o in first.context["month_options"]] == ["2026-08", "2026-07"]
    june = client.get(reverse("employee_list"), {"oy": "2026-06"})
    assert june.context["davomad_url"] == f"{reverse('attendance_grid')}?oy=2026-07"


def test_marking_an_absence_writes_a_row_and_clearing_it_removes_the_row(
    client, admin_user, on_august_12
):
    worker = make_worker()
    client.force_login(admin_user)
    post_sheet(client, worker, {("d", "2026-08-11"): "kelmadi"})
    row = Attendance.objects.get(employee=worker)
    assert (row.date, row.present) == (date(2026, 8, 11), False)

    post_sheet(client, worker, {("d", "2026-08-11"): "keldi"})
    assert not Attendance.objects.filter(employee=worker).exists()


def test_a_closed_day_only_moves_when_it_was_opened_on_purpose(
    client, admin_user, on_august_12
):
    worker = make_worker()
    client.force_login(admin_user)
    post_sheet(client, worker, {("d", "2026-08-10"): "kelmadi"})
    assert not Attendance.objects.filter(employee=worker).exists()

    post_sheet(client, worker, {("d", "2026-08-10"): "kelmadi", ("o", "2026-08-10"): "1"})
    assert Attendance.objects.filter(employee=worker, date=date(2026, 8, 10)).exists()


def test_a_future_day_cannot_be_marked(client, admin_user, on_august_12):
    worker = make_worker()
    client.force_login(admin_user)
    post_sheet(client, worker, {("d", "2026-08-13"): "kelmadi"})
    assert not Attendance.objects.filter(employee=worker).exists()


def test_the_sheet_lists_only_accounts_open_that_month(client, admin_user, on_august_12):
    here = make_worker(name="Shu oyda bor")
    later = make_worker(name="Keyin keladi", start=date(2026, 9, 1))
    client.force_login(admin_user)
    response = client.get(reverse("attendance_grid"), {"oy": "2026-08"})
    listed = {r["employee"].pk for r in response.context["rows"]}
    assert listed == {here.pk}
    assert later.pk not in listed


def test_the_sheet_exports(client, admin_user, on_august_12):
    worker = make_worker()
    mark(worker, MONDAY, present=False)
    client.force_login(admin_user)
    response = client.get(reverse("attendance_excel"), {"oy": "2026-07"})
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ["Davomad", "Oylik", "Izohlar", "Bayramlar"]
    wages = list(book["Oylik"].iter_rows(values_only=True))
    assert wages[0][5] == "Hisoblangan"
    assert (wages[1][0], wages[1][5]) == (worker.name, float(WAGE - DAY))


# --- kun kartasi -----------------------------------------------------------------------


def test_the_day_card_books_leave_ahead_and_records_what_happened(
    client, admin_user, on_august_12
):
    worker = make_worker()
    client.force_login(admin_user)
    ahead = reverse("employee_day", args=[worker.pk, "2026-08-20"])
    client.post(ahead, {"action": "dam_haqli", "note": "to'y"})
    leave = Leave.objects.get(employee=worker)
    assert (leave.date, leave.paid, leave.note) == (date(2026, 8, 20), True, "to'y")

    # A day long closed on the sheet is still correctable from its own card.
    behind = reverse("employee_day", args=[worker.pk, "2026-08-03"])
    client.post(behind, {"action": "kelmadi", "note": "kasal"})
    row = Attendance.objects.get(employee=worker)
    assert (row.date, row.present, row.note) == (date(2026, 8, 3), False, "kasal")


# --- ish vaqti -------------------------------------------------------------------------


def test_the_day_card_changes_one_days_hours_and_moves_no_money(
    client, admin_user, on_august_12
):
    """Hours say WHEN a day is worked, not whether: a short day is still a day."""
    worker = make_worker(start=date(2026, 8, 1))
    assert worker.work_hours == "09:00–19:00"               # what a new worker gets
    client.force_login(admin_user)
    wage = worker.earned_in(2026, 8)
    card = reverse("employee_day", args=[worker.pk, "2026-08-20"])

    client.post(card, {
        "action": "ish_soati", "work_start": "10:00", "work_end": "15:00",
        "note": "shifokor",
    })
    row = DayHours.objects.get(employee=worker)
    assert (row.date, row.label, row.note) == (date(2026, 8, 20), "10:00–15:00", "shifokor")
    assert worker.earned_in(2026, 8) == wage
    assert not Attendance.objects.exists() and not Leave.objects.exists()

    # Half a working day is not an answer, so nothing is overwritten with it.
    client.post(card, {"action": "ish_soati", "work_start": "11:00", "work_end": ""})
    assert DayHours.objects.get(employee=worker).label == "10:00–15:00"

    # The usual hours put back remove the exception rather than repeat them.
    client.post(card, {"action": "ish_soati", "work_start": "09:00", "work_end": "19:00"})
    assert not DayHours.objects.exists()

    # A Sunday nobody works has no hours to change, so its card does not ask.
    sunday = client.get(reverse("employee_day", args=[worker.pk, "2026-08-16"]))
    assert "ish_soati" in client.get(card).content.decode()
    assert "ish_soati" not in sunday.content.decode()


def test_the_calendar_colours_a_day_by_what_happened_and_prints_its_hours(
    client, admin_user, on_august_12
):
    worker = make_worker(start=date(2026, 8, 1))
    mark(worker, date(2026, 8, 3), present=False)           # a Monday missed
    mark(worker, date(2026, 8, 9))                          # a Sunday worked
    Holiday.objects.create(date=date(2026, 8, 25), name="Hayit", rest=True)
    DayHours.objects.create(
        employee=worker, date=date(2026, 8, 20), start=time(10), end=time(15)
    )
    client.force_login(admin_user)
    weeks = client.get(reverse("employee_detail", args=[worker.pk])).context["weeks"]
    cells = {c["date"].day: c for row in weeks for c in row if not c["outside"]}

    #       Sunday  missed    worked     Sunday worked  ahead     ahead, own hours  holiday
    days = [2,      3,        4,         9,             13,       20,               25]
    assert [cells[d]["tone"] for d in days] == [
        "rest", "absent", "present", "present", "future", "future", "holiday",
    ]
    assert [cells[d]["hours"] for d in days] == [
        "", "", "09:00–19:00", "09:00–19:00", "09:00–19:00", "10:00–15:00", "",
    ]
    assert [d for d in days if cells[d]["hours_custom"]] == [20]


def test_a_workers_hours_are_both_ends_or_neither():
    month = timezone.localdate().strftime("%Y-%m")
    data = {"name": "Ишчи Аваз", "salary": "2700000", "start_month": month,
            "opening_balance": "0"}
    assert EmployeeForm({**data, "work_start": "", "work_end": ""}).is_valid()
    assert EmployeeForm({**data, "work_start": "20:00", "work_end": "08:00"}).is_valid()
    half = EmployeeForm({**data, "work_start": "09:00", "work_end": ""})
    assert not half.is_valid() and "work_end" in half.errors


def test_a_row_of_the_workers_page_edits_that_row_and_nothing_else(client, admin_user):
    """The pencils on a worker's page open the full form cut down to one row, so what
    is not on the card is not touched — and a wage changed there is still dated."""
    worker = make_worker()
    client.force_login(admin_user)
    hours = reverse("employee_edit_field", args=[worker.pk, "ish-vaqti"])
    card = client.get(hours, headers={"x-requested-with": "XMLHttpRequest"})
    assert 'name="work_start"' in card.content.decode()
    assert 'name="salary"' not in card.content.decode()

    client.post(hours, {"work_start": "08:00", "work_end": "17:00"})
    worker.refresh_from_db()
    assert worker.work_hours == "08:00–17:00"
    assert (worker.name, worker.salary, worker.is_active) == ("Ишчи Аваз", WAGE, True)

    wage = reverse("employee_edit_field", args=[worker.pk, "oylik"])
    this_month = timezone.localdate().replace(day=1)
    client.post(wage, {"salary": "3000000", "salary_from": this_month.strftime("%Y-%m")})
    worker.refresh_from_db()
    assert (worker.salary, worker.work_hours) == (Decimal("3000000"), "08:00–17:00")
    assert SalaryRate.objects.filter(
        employee=worker, amount=Decimal("3000000"), effective_from=this_month
    ).exists()

    # A raise still cannot start before the account does, though the account's start
    # is not on this card to compare against.
    early = client.post(
        wage, {"salary": "4000000", "salary_from": "2026-06"},
        headers={"x-requested-with": "XMLHttpRequest"},
    )
    worker.refresh_from_db()
    assert (early.status_code, worker.salary) == (422, Decimal("3000000"))

    unknown = reverse("employee_edit_field", args=[worker.pk, "ism"])
    assert client.get(unknown).status_code == 404


# --- bayramlar -------------------------------------------------------------------------


def test_fixed_holidays_are_seeded_and_a_year_can_be_filled_in(client, admin_user):
    assert Holiday.objects.filter(date=date(2026, 9, 1), rest=True).exists()
    client.force_login(admin_user)
    client.post(f"{reverse('holiday_seed')}?yil=2035")
    assert Holiday.objects.filter(date__year=2035).count() == len(Holiday.FIXED)

    client.post(reverse("holiday_create"), {
        "date": "2027-03-10", "name": "Ramazon hayit", "rest": "on", "note": "",
    })
    assert Holiday.objects.filter(date=date(2027, 3, 10), rest=True).exists()


# --- who may do what ---------------------------------------------------------------------


def test_a_seller_reads_the_payroll_but_cannot_move_a_wage(client, payroll_seller, on_august_12):
    worker = make_worker()
    holiday = Holiday.objects.first()
    client.force_login(payroll_seller)

    assert client.get(reverse("employee_list")).status_code == 200
    detail = client.get(reverse("employee_detail", args=[worker.pk]))
    assert detail.status_code == 200
    day_url = reverse("employee_day", args=[worker.pk, "2026-08-11"])
    assert day_url not in detail.content.decode()           # no pencil on the calendar

    forbidden = [
        client.get(reverse("attendance_grid")),
        client.post(f"{reverse('attendance_grid')}?oy=2026-08", {
            f"d-{worker.pk}-2026-08-11": "kelmadi",
        }),
        client.get(reverse("attendance_excel")),
        client.post(day_url, {"action": "kelmadi", "note": ""}),
        client.get(reverse("holiday_list")),
        client.post(reverse("holiday_create"), {"date": "2027-03-10", "name": "X"}),
        client.post(reverse("holiday_delete", args=[holiday.pk])),
        client.post(f"{reverse('holiday_seed')}?yil=2035"),
    ]
    assert [r.status_code for r in forbidden] == [403] * len(forbidden)
    assert not Attendance.objects.exists()
    assert Holiday.objects.filter(pk=holiday.pk).exists()

def test_a_seller_without_payroll_access_has_no_way_in(client, seller_user, on_august_12):
    """«Xodimlar oyligi» is opened account by account: a seller it was not switched on
    for sees no menu entry and is refused on every page behind it."""
    worker = make_worker()
    client.force_login(seller_user)

    home = client.get(reverse("kassa"))
    assert home.status_code == 200
    assert reverse("employee_list") not in home.content.decode()

    forbidden = [
        client.get(reverse("employee_list")),
        client.get(reverse("employee_export")),
        client.get(reverse("employee_detail", args=[worker.pk])),
        client.post(reverse("employee_create"), {
            "name": "X", "salary": "1000000", "start_month": "2026-08",
            "is_active": "on",
        }),
        client.post(reverse("employee_edit", args=[worker.pk]), {"name": "Y"}),
        client.post(reverse("employee_edit_field", args=[worker.pk, "oylik"]), {
            "salary": "1",
        }),
        client.post(reverse("employee_delete", args=[worker.pk])),
    ]
    assert [r.status_code for r in forbidden] == [403] * len(forbidden)
    worker.refresh_from_db()
    assert (worker.name, worker.salary) == ("Ишчи Аваз", WAGE)
    assert Employee.objects.count() == 1


def test_switching_payroll_access_on_opens_the_menu_and_the_page(client, seller_user, admin_user):
    assert not seller_user.can_see_payroll
    assert admin_user.can_see_payroll                       # no flag needed for an admin
    seller_user.payroll_access = True
    seller_user.save(update_fields=["payroll_access"])
    client.force_login(seller_user)

    page = client.get(reverse("employee_list"))
    assert page.status_code == 200
    assert f'href="{reverse("employee_list")}"' in page.content.decode()


def test_the_hr_pages_open_for_an_admin(client, admin_user, on_august_12):
    worker = make_worker()
    other = make_worker(name="Ишчи Собир")
    mark(worker, date(2026, 8, 3), present=False)
    Leave.objects.create(employee=worker, date=date(2026, 8, 20), paid=False)
    client.force_login(admin_user)

    pages = [
        reverse("employee_list"),
        f"{reverse('employee_list')}?q=Собир",
        reverse("attendance_grid"),
        f"{reverse('attendance_grid')}?oy=2026-07",
        reverse("employee_detail", args=[worker.pk]),
        f"{reverse('employee_detail', args=[other.pk])}?oy=2026-07",
        reverse("holiday_list"),
        reverse("employee_day", args=[worker.pk, "2026-08-11"]),
        reverse("employee_day", args=[worker.pk, "2026-08-20"]),
        reverse("employee_export"),
    ]
    assert [client.get(url).status_code for url in pages] == [200] * len(pages)

    detail = client.get(reverse("employee_detail", args=[worker.pk])).content.decode()
    assert reverse("employee_day", args=[worker.pk, "2026-08-11"]) in detail
    # The sheet follows the same search as the payroll page.
    found = client.get(f"{reverse('attendance_grid')}?q=Собир")
    assert [r["employee"].pk for r in found.context["rows"]] == [other.pk]


def test_sunday_rests_everybody():
    """One week for everyone: six days on, Sunday off. There is no second rota."""
    worker = make_worker()
    assert not hasattr(worker, "work_mode")
    assert SUNDAY in worker.rest_dates_in(YEAR, MONTH)
    mark(worker, SUNDAY, present=False)                     # resting on it costs nothing
    assert worker.earned_in(YEAR, MONTH) == WAGE
