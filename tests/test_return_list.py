"""Vazvratlar — every return on one page, with where its value went.

Pins down what the page promises: the value of each return is split into the part
that cancelled debt and the part handed back (advance or cash), the totals add back
up, a seller sees only the returns on their own sales, and the Excel holds the same
rows the page shows.
"""

from decimal import Decimal

import pytest
from django.utils import timezone

from accounts.models import User
from crm.models import Client, Payment, Return, Sale, SaleItem

pytestmark = pytest.mark.django_db


def _return(sale, weight, settlement_amount=None, kind=Payment.Kind.RETURN_CREDIT):
    item = sale.items.first()
    ret = Return.objects.create(
        sale_item=item, weight=Decimal(weight), created_by=sale.sales_rep,
        date=timezone.localdate(),
    )
    if settlement_amount is not None:
        ret.settlement = Payment.objects.create(
            sale=sale, client=sale.client, amount=Decimal(settlement_amount),
            method=Payment.Method.CASH, kind=kind, date=ret.date,
            created_by=sale.sales_rep,
        )
        ret.save(update_fields=["settlement"])
    return ret


def test_return_value_splits_into_debt_and_advance(client, admin_user, sample_data):
    """Sale 150 000, paid 50 000. Returning 2 kg (30 000) only shrinks the debt; a
    further 8 kg (120 000) cancels the remaining 70 000 of debt and parks 50 000 as
    advance."""
    sale = sample_data["sale"]
    _return(sale, "2")
    _return(sale, "8", settlement_amount="50000")

    client.force_login(admin_user)
    resp = client.get("/vazvrat/")
    assert resp.status_code == 200
    totals = resp.context["totals"]
    assert totals["total"] == Decimal("150000")
    assert totals["to_debt"] == Decimal("100000")
    assert totals["advance"] == Decimal("50000")
    assert totals["refund"] == Decimal("0")
    assert totals["count"] == 2 and totals["clients"] == 1
    assert totals["weight_kg"] == Decimal("10")

    rows = {row["ret"].weight: row for row in resp.context["page"]}
    assert rows[Decimal("2")]["to_debt"] == Decimal("30000")
    assert rows[Decimal("2")]["back"] == Decimal("0")
    assert rows[Decimal("8")]["to_debt"] == Decimal("70000")
    assert rows[Decimal("8")]["back"] == Decimal("50000")
    assert not rows[Decimal("8")]["back_is_refund"]
    assert "Test mijoz" in resp.content.decode()


def test_cash_refund_is_counted_apart_from_advance(client, admin_user, sample_data):
    _return(sample_data["sale"], "10", settlement_amount="50000",
            kind=Payment.Kind.REFUND_OUT)
    client.force_login(admin_user)
    resp = client.get("/vazvrat/")
    totals = resp.context["totals"]
    assert totals["refund"] == Decimal("50000")
    assert totals["advance"] == Decimal("0")
    assert list(resp.context["page"])[0]["back_is_refund"]


def test_settle_filter_picks_where_the_money_went(client, admin_user, sample_data):
    sale = sample_data["sale"]
    _return(sale, "2")
    _return(sale, "8", settlement_amount="50000")
    client.force_login(admin_user)

    only_debt = client.get("/vazvrat/", {"settle": "debt"}).context["page"]
    assert [row["ret"].weight for row in only_debt] == [Decimal("2")]
    advance = client.get("/vazvrat/", {"settle": "advance"}).context["page"]
    assert [row["ret"].weight for row in advance] == [Decimal("8")]
    assert not list(client.get("/vazvrat/", {"settle": "refund"}).context["page"])


def test_search_finds_a_return_by_product_name(client, admin_user, sample_data):
    _return(sample_data["sale"], "2")
    client.force_login(admin_user)
    assert len(list(client.get("/vazvrat/", {"q": "paket"}).context["page"])) == 1
    assert not list(client.get("/vazvrat/", {"q": "yo'q-narsa"}).context["page"])


def test_seller_sees_only_returns_on_their_own_sales(
    client, admin_user, seller_user, sample_data
):
    other = User.objects.create_user(
        username="boshqa", password="x", role=User.Role.SALES,
    )
    other_client = Client.objects.create(name="Boshqa mijoz", owner=other)
    other_sale = Sale.objects.create(client=other_client, sales_rep=other)
    SaleItem.objects.create(
        sale=other_sale, product=sample_data["product"], dimension=Sale.Dimension.KG,
        weight=Decimal("5"), price=Decimal("15000"), cost_price=Decimal("10000"),
    )
    _return(sample_data["sale"], "1")
    _return(other_sale, "1")

    client.force_login(seller_user)
    resp = client.get("/vazvrat/")
    assert [row["ret"].sale_id for row in resp.context["page"]] == [sample_data["sale"].pk]
    assert "Boshqa mijoz" not in resp.content.decode()

    client.force_login(admin_user)
    assert len(list(client.get("/vazvrat/").context["page"])) == 2


def test_export_holds_the_filtered_rows(client, admin_user, sample_data):
    from io import BytesIO

    from openpyxl import load_workbook

    sale = sample_data["sale"]
    _return(sale, "2")
    _return(sale, "8", settlement_amount="50000")
    client.force_login(admin_user)
    resp = client.get("/vazvrat/export/", {"settle": "advance"})
    assert resp.status_code == 200
    ws = load_workbook(BytesIO(resp.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][8:11] == ("Qarzdan ayirildi", "Avansga o'tdi", "Naqd qaytarildi")
    assert len(rows) == 2
    assert rows[1][7] == 120000 and rows[1][8] == 70000 and rows[1][9] == 50000
