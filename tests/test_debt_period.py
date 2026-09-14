"""Qarzlar — the sana oralig'i beside the Excel button.

The page is a balance list: it says who owes what today. Picking a window turns it
into a period sverka as well — "from the 14th to the 14th, this client took this much
goods and handed over this much money" — without disturbing the balances beside it.

What is pinned down here: the window only shows up when it is asked for, the two
figures mean what the page says they mean (returns netted off, advance credit being
spent not counted twice), a client who settled up inside the window is still on the
report, and the Excel carries the same columns with the period named in them.
"""

from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from crm.models import Client, Payment, Product, Return, Sale, SaleItem

pytestmark = pytest.mark.django_db

TODAY = timezone.localdate()


def _sale(client, seller, *, day, kg, price, product, paid=None, term=7):
    sale = Sale.objects.create(
        client=client, sales_rep=seller, date=day,
        debt_deadline=day + timedelta(days=term), debt_term_days=term,
    )
    SaleItem.objects.create(
        sale=sale, product=product, dimension=Sale.Dimension.KG,
        weight=Decimal(kg), price=Decimal(price), cost_price=Decimal("10000"),
    )
    if paid:
        Payment.objects.create(
            sale=sale, amount=Decimal(paid), amount_original=Decimal(paid),
            method=Payment.Method.CASH, kind=Payment.Kind.SALE,
            date=day, created_by=seller,
        )
    return sale


@pytest.fixture
def product(db):
    return Product.objects.create(
        name="Sverka paket", sku="SVR-001",
        cost_price=Decimal("10000"), price=Decimal("15000"),
    )


def _window(dan, gacha):
    return {"dan": dan.isoformat(), "gacha": gacha.isoformat()}


def test_no_window_leaves_the_page_exactly_as_it_was(client, admin_user, sample_data):
    """A debtors list has no period of its own, so nothing is dated until dates are
    picked: no sverka columns, no extra rows."""
    client.force_login(admin_user)
    resp = client.get("/debts/")

    assert resp.status_code == 200
    assert resp.context["period_window"] is False
    assert resp.context["is_all"] is True
    assert all(row["period"] is None for row in resp.context["debtors"])
    assert "Oraliqda olgan yuk" not in resp.content.decode()


def test_window_shows_goods_taken_and_money_paid(
    client, admin_user, seller_user, product
):
    """The client's example: 14th to 14th. Two sales inside the window and one before
    it; the earlier one is part of the balance but not of the period figures."""
    mijoz = Client.objects.create(name="Oraliq mijoz", owner=seller_user)
    dan, gacha = TODAY - timedelta(days=31), TODAY - timedelta(days=1)
    _sale(mijoz, seller_user, day=dan - timedelta(days=10), kg="5", price="15000",
          product=product)                                        # before the window
    _sale(mijoz, seller_user, day=dan + timedelta(days=2), kg="10", price="15000",
          product=product, paid="60000")                          # inside
    _sale(mijoz, seller_user, day=gacha, kg="4", price="20000", product=product)

    client.force_login(admin_user)
    resp = client.get("/debts/", _window(dan, gacha))
    row = next(r for r in resp.context["debtors"] if r["client"] == mijoz)

    assert resp.context["period_window"] is True
    # 10 kg x 15 000 + 4 kg x 20 000 — the pre-window sale stays out.
    assert row["period"]["net_taken"] == Decimal("230000")
    assert row["period"]["net_taken_kg"] == Decimal("14")
    assert row["period"]["paid"] == Decimal("60000")
    assert row["period"]["pay_count"] == 1
    # The balance beside it is still every open receipt, window or no window.
    assert row["remaining"] == Decimal("75000") + Decimal("90000") + Decimal("80000")


def test_returned_goods_come_off_what_was_taken(
    client, admin_user, seller_user, product
):
    """Goods that came back are goods the client did not keep — but the returned
    figure stays visible beside the net one instead of vanishing into it."""
    mijoz = Client.objects.create(name="Qaytargan mijoz", owner=seller_user)
    dan, gacha = TODAY - timedelta(days=10), TODAY
    sale = _sale(mijoz, seller_user, day=dan, kg="10", price="15000", product=product)
    Return.objects.create(
        sale=sale, sale_item=sale.items.first(), product=product,
        dimension=Sale.Dimension.KG, weight=Decimal("3"), date=gacha,
        restock=True, created_by=seller_user,
    )

    client.force_login(admin_user)
    resp = client.get("/debts/", _window(dan, gacha))
    row = next(r for r in resp.context["debtors"] if r["client"] == mijoz)

    assert row["period"]["taken"] == Decimal("150000")
    assert row["period"]["returned"] == Decimal("45000")
    assert row["period"]["net_taken"] == Decimal("105000")
    assert row["period"]["net_taken_kg"] == Decimal("7")


def test_spending_advance_credit_is_not_counted_as_money_paid(
    client, admin_user, seller_user, product
):
    """The same rule the to'lovlar page runs on: the deposit is the money the client
    handed over, and closing a receipt out of that credit is not a second payment."""
    mijoz = Client.objects.create(name="Avansli mijoz", owner=seller_user)
    dan, gacha = TODAY - timedelta(days=5), TODAY
    sale = _sale(mijoz, seller_user, day=dan, kg="10", price="15000", product=product)
    Payment.objects.create(
        client=mijoz, amount=Decimal("100000"), amount_original=Decimal("100000"),
        method=Payment.Method.CASH, kind=Payment.Kind.ADVANCE_IN,
        date=dan, created_by=seller_user,
    )
    Payment.objects.create(
        sale=sale, amount=Decimal("100000"), amount_original=Decimal("100000"),
        method=Payment.Method.CASH, kind=Payment.Kind.ADVANCE_USED,
        date=gacha, created_by=seller_user,
    )

    client.force_login(admin_user)
    resp = client.get("/debts/", _window(dan, gacha))
    row = next(r for r in resp.context["debtors"] if r["client"] == mijoz)

    assert row["period"]["paid"] == Decimal("100000")
    assert row["period"]["pay_count"] == 1


def test_client_who_settled_up_inside_the_window_still_shows(
    client, admin_user, seller_user, product
):
    """They owe nothing and hold no advance, so no balance puts them on the page — and
    they are exactly the client a period report must not lose. The Qarzdorlar switch
    filters them straight back out, which is right: they are not debtors."""
    mijoz = Client.objects.create(name="To'lab ketgan", owner=seller_user)
    dan, gacha = TODAY - timedelta(days=6), TODAY
    _sale(mijoz, seller_user, day=dan, kg="10", price="15000", product=product,
          paid="150000")

    client.force_login(admin_user)
    resp = client.get("/debts/", _window(dan, gacha))
    row = next((r for r in resp.context["debtors"] if r["client"] == mijoz), None)

    assert row is not None
    assert row["remaining"] == Decimal("0")
    assert row["period"]["net_taken"] == Decimal("150000")
    assert row["period"]["paid"] == Decimal("150000")

    only_debtors = client.get("/debts/", {**_window(dan, gacha), "tur": "qarz"})
    assert mijoz not in [r["client"] for r in only_debtors.context["debtors"]]

    # The balance cards are a today figure and must not move because a window was
    # picked: this client owes nothing and adds nothing to them.
    assert resp.context["total_debt"] == Decimal("0")
    assert resp.context["total_debtors"] == 0


def test_period_totals_add_up_the_rows_on_the_page(
    client, admin_user, seller_user, product
):
    dan, gacha = TODAY - timedelta(days=4), TODAY
    for name, kg, paid in (("Bir", "10", "50000"), ("Ikki", "6", "20000")):
        mijoz = Client.objects.create(name=name, owner=seller_user)
        _sale(mijoz, seller_user, day=dan, kg=kg, price="15000", product=product,
              paid=paid)

    client.force_login(admin_user)
    resp = client.get("/debts/", _window(dan, gacha))
    totals = resp.context["totals"]

    assert totals["taken"] == Decimal("240000")   # (10 + 6) kg x 15 000
    assert totals["taken_kg"] == Decimal("16")
    assert totals["paid"] == Decimal("70000")


def test_seller_sees_only_their_own_period_figures(
    django_user_model, client, seller_user, product
):
    """Same scoping as every other figure on the page: a seller's sverka is their own
    goods and their own money, never the other seller's."""
    other = django_user_model.objects.create_user(
        username="boshqa_sotuvchi", password="x", role=django_user_model.Role.SALES,
    )
    mijoz = Client.objects.create(name="Umumiy mijoz", owner=seller_user)
    dan, gacha = TODAY - timedelta(days=3), TODAY
    _sale(mijoz, seller_user, day=dan, kg="10", price="15000", product=product,
          paid="30000")
    _sale(mijoz, other, day=dan, kg="8", price="15000", product=product, paid="90000")

    client.force_login(seller_user)
    resp = client.get("/debts/", _window(dan, gacha))
    row = next(r for r in resp.context["debtors"] if r["client"] == mijoz)

    assert row["period"]["net_taken"] == Decimal("150000")
    assert row["period"]["paid"] == Decimal("30000")


def test_excel_carries_the_window_in_its_headers(
    client, admin_user, seller_user, product
):
    """The file is passed around on its own, so a column of period figures that does
    not say which period it belongs to would be worse than no column."""
    mijoz = Client.objects.create(name="Excel mijoz", owner=seller_user)
    dan, gacha = TODAY - timedelta(days=8), TODAY
    _sale(mijoz, seller_user, day=dan, kg="10", price="15000", product=product,
          paid="40000")

    client.force_login(admin_user)
    resp = client.get("/debts/export/", _window(dan, gacha))
    sheet = load_workbook(BytesIO(resp.content)).active
    headers = [c.value for c in sheet[1]]
    row = dict(zip(headers, [c.value for c in sheet[2]]))
    span = f"({dan.strftime('%d.%m.%Y')}–{gacha.strftime('%d.%m.%Y')})"

    assert f"Olgan yuk, so'm {span}" in headers
    assert row[f"Olgan yuk, so'm {span}"] == 150000
    assert row[f"Olgan yuk, kg {span}"] == 10
    assert row[f"To'lagan pul {span}"] == 40000
    assert row["Qarz qoldig'i"] == 110000


def test_excel_without_a_window_keeps_its_old_columns(client, admin_user, sample_data):
    client.force_login(admin_user)
    resp = client.get("/debts/export/")
    headers = [c.value for c in load_workbook(BytesIO(resp.content)).active[1]]

    assert headers[-1] == "Avans qachondan"
    assert not any("Olgan yuk" in (h or "") for h in headers)
